"""严格基于官方模板导出问题二最终结果并进行回读校验。

分类结果只取自现有 q2_keyword_scores.csv 的“最终类别”，本脚本不计算得分、
不训练模型、不更改分类。方案、推广单元、序号和关键词以 Attachment1.xlsx
Sheet3 的原始单元格值写出，避免 CSV/Excel 类型转换改变编码。
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCORES_CSV = PROJECT_ROOT / "output" / "tables" / "q2_keyword_scores.csv"
CATEGORY_COUNTS_CSV = PROJECT_ROOT / "output" / "tables" / "q2_category_counts.csv"
SOURCE_XLSX = PROJECT_ROOT / "data" / "raw" / "Attachment" / "Attachment1.xlsx"
TEMPLATE_XLSX = (
    PROJECT_ROOT / "data" / "raw" / "Attachment" / "Attachment2" / "result2.xlsx"
)
OUTPUT_XLSX = PROJECT_ROOT / "output" / "result2.xlsx"
VALIDATION_CSV = PROJECT_ROOT / "output" / "tables" / "q2_result2_validation.csv"

SOURCE_SHEET = "Sheet3"
SCORE_COLUMNS = ["序号", "关键词", "方案ID", "推广单元ID", "最终类别"]
CATEGORY_ORDER = ["黄金词", "重点词", "潜力词", "问题词", "无效词"]
EXPECTED_COUNTS = {
    "黄金词": 198,
    "重点词": 502,
    "潜力词": 461,
    "问题词": 178,
    "无效词": 888,
}
EXPECTED_TOTAL = 2227
EXPECTED_TEMPLATE_HEADERS = [
    "方案ID",
    "推广单元",
    "序号",
    "黄金词",
    "重点词",
    "潜力词",
    "问题词",
    "无效词",
]


def display_value(value: Any) -> str | int | float | bool:
    """把校验值转换成适合写入 CSV 的稳定标量。"""
    if isinstance(value, (list, tuple)):
        return " | ".join(str(item) for item in value)
    if isinstance(value, dict):
        return " | ".join(f"{key}:{item}" for key, item in value.items())
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, (np.bool_,)):
        return bool(value)
    return value


def add_check(
    checks: list[dict[str, Any]],
    check_name: str,
    expected: Any,
    actual: Any,
    passed: bool,
) -> None:
    """追加一条统一格式的校验记录。"""
    checks.append(
        {
            "check_name": check_name,
            "expected": display_value(expected),
            "actual": display_value(actual),
            "passed": bool(passed),
        }
    )


def save_validation(checks: list[dict[str, Any]]) -> pd.DataFrame:
    """保存当前已完成的校验；失败时也留下明确证据。"""
    validation = pd.DataFrame(
        checks,
        columns=["check_name", "expected", "actual", "passed"],
    )
    VALIDATION_CSV.parent.mkdir(parents=True, exist_ok=True)
    validation.to_csv(VALIDATION_CSV, index=False, encoding="utf-8-sig")
    return validation


def require_all_passed(checks: list[dict[str, Any]], stage: str) -> None:
    """任一校验失败即保存校验表并明确停止。"""
    failed = [row for row in checks if not row["passed"]]
    if failed:
        save_validation(checks)
        names = "、".join(str(row["check_name"]) for row in failed)
        raise RuntimeError(f"{stage}失败，已停止生成 result2.xlsx：{names}")


def canonical_code(value: Any) -> str | None:
    """无损比较 Excel/CSV 编码；字符串保留可能存在的前导零。"""
    if value is None or pd.isna(value):
        return None
    if isinstance(value, (bool, np.bool_)):
        return str(bool(value))
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)):
        numeric = float(value)
        if not np.isfinite(numeric):
            return None
        return str(int(numeric)) if numeric.is_integer() else format(numeric, ".15g")
    text = str(value).strip()
    return text if text != "" else None


def parse_sequence(value: Any) -> int | None:
    """严格解析整数序号。"""
    text = canonical_code(value)
    if text is None:
        return None
    try:
        numeric = float(text)
    except ValueError:
        return None
    if not np.isfinite(numeric) or not numeric.is_integer():
        return None
    return int(numeric)


def is_nonempty(value: Any) -> bool:
    """判断结果单元格是否真正含值；数值0也视为一个实际值。"""
    if value is None:
        return False
    return not (isinstance(value, str) and value.strip() == "")


def load_scores(checks: list[dict[str, Any]]) -> pd.DataFrame:
    """读取定稿分类，并执行全部生成前分类一致性检查。"""
    if not SCORES_CSV.exists():
        raise FileNotFoundError(f"未找到定稿分类表：{SCORES_CSV}")
    scores = pd.read_csv(
        SCORES_CSV,
        usecols=lambda column: column in SCORE_COLUMNS,
        dtype=str,
        keep_default_na=False,
    )
    missing_columns = sorted(set(SCORE_COLUMNS).difference(scores.columns))
    add_check(
        checks,
        "input_required_columns",
        SCORE_COLUMNS,
        scores.columns.tolist(),
        not missing_columns,
    )
    if missing_columns:
        require_all_passed(checks, "评分表字段检查")

    scores = scores[SCORE_COLUMNS].copy()
    scores["序号_整数"] = scores["序号"].map(parse_sequence)
    valid_sequence_count = int(scores["序号_整数"].notna().sum())
    unique_sequence_count = int(scores["序号_整数"].nunique(dropna=True))
    missing_category_count = int(scores["最终类别"].map(canonical_code).isna().sum())
    actual_categories = sorted(set(scores["最终类别"]) - {""})

    add_check(checks, "input_record_count", EXPECTED_TOTAL, len(scores), len(scores) == EXPECTED_TOTAL)
    add_check(
        checks,
        "input_sequence_nonmissing_count",
        EXPECTED_TOTAL,
        valid_sequence_count,
        valid_sequence_count == EXPECTED_TOTAL,
    )
    add_check(
        checks,
        "input_sequence_unique_count",
        EXPECTED_TOTAL,
        unique_sequence_count,
        unique_sequence_count == EXPECTED_TOTAL,
    )
    add_check(
        checks,
        "input_missing_category_count",
        0,
        missing_category_count,
        missing_category_count == 0,
    )
    add_check(
        checks,
        "input_allowed_categories",
        CATEGORY_ORDER,
        actual_categories,
        set(actual_categories) == set(CATEGORY_ORDER),
    )

    actual_counts = scores["最终类别"].value_counts().to_dict()
    for category in CATEGORY_ORDER:
        actual = int(actual_counts.get(category, 0))
        expected = EXPECTED_COUNTS[category]
        add_check(
            checks,
            f"input_category_count_{category}",
            expected,
            actual,
            actual == expected,
        )
    category_total = int(sum(actual_counts.get(category, 0) for category in CATEGORY_ORDER))
    add_check(
        checks,
        "input_category_count_total",
        EXPECTED_TOTAL,
        category_total,
        category_total == EXPECTED_TOTAL,
    )

    if CATEGORY_COUNTS_CSV.exists():
        category_counts = pd.read_csv(CATEGORY_COUNTS_CSV, dtype=str, keep_default_na=False)
        required = {"最终类别", "数量"}
        counts_match = required.issubset(category_counts.columns)
        counts_from_file: dict[str, int] = {}
        if counts_match:
            try:
                counts_from_file = {
                    str(row["最终类别"]): int(row["数量"])
                    for _, row in category_counts.iterrows()
                }
            except (TypeError, ValueError):
                counts_match = False
        counts_match = counts_match and all(
            counts_from_file.get(category) == EXPECTED_COUNTS[category]
            for category in CATEGORY_ORDER
        )
        add_check(
            checks,
            "category_counts_file_consistency",
            EXPECTED_COUNTS,
            counts_from_file,
            counts_match,
        )

    require_all_passed(checks, "生成前分类一致性检查")
    return scores.sort_values("序号_整数", kind="stable").reset_index(drop=True)


def load_source_rows(checks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """从 Attachment1 Sheet3 读取原始标识值及其 Excel 数据类型。"""
    if not SOURCE_XLSX.exists():
        raise FileNotFoundError(f"未找到原始数据：{SOURCE_XLSX}")
    workbook = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    add_check(
        checks,
        "source_sheet_exists",
        SOURCE_SHEET,
        workbook.sheetnames,
        SOURCE_SHEET in workbook.sheetnames,
    )
    require_all_passed(checks, "原始数据工作表检查")
    sheet = workbook[SOURCE_SHEET]
    headers = [cell.value for cell in sheet[1]]
    header_map = {str(value).strip(): index + 1 for index, value in enumerate(headers)}
    required_headers = ["序号", "关键词", "方案ID", "推广单元ID"]
    add_check(
        checks,
        "source_required_columns",
        required_headers,
        [header for header in required_headers if header in header_map],
        all(header in header_map for header in required_headers),
    )
    require_all_passed(checks, "原始数据字段检查")

    rows: list[dict[str, Any]] = []
    incomplete_row_count = 0
    # read_only 工作表必须顺序迭代；反复调用 sheet.cell 会从文件头重复扫描。
    for row_number, row_values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        values = {
            header: row_values[header_map[header] - 1]
            for header in required_headers
        }
        if all(value is None for value in values.values()):
            continue
        if any(canonical_code(value) is None for value in values.values()):
            incomplete_row_count += 1
        values["序号_整数"] = parse_sequence(values["序号"])
        values["源行号"] = row_number
        rows.append(values)
    workbook.close()

    source_sequences = [row["序号_整数"] for row in rows]
    unique_count = len({value for value in source_sequences if value is not None})
    add_check(checks, "source_record_count", EXPECTED_TOTAL, len(rows), len(rows) == EXPECTED_TOTAL)
    add_check(
        checks,
        "source_incomplete_identifier_rows",
        0,
        incomplete_row_count,
        incomplete_row_count == 0,
    )
    add_check(
        checks,
        "source_sequence_unique_count",
        EXPECTED_TOTAL,
        unique_count,
        unique_count == EXPECTED_TOTAL,
    )
    require_all_passed(checks, "原始数据记录检查")
    return sorted(rows, key=lambda row: int(row["序号_整数"]))


def validate_score_source_mapping(
    scores: pd.DataFrame,
    source_rows: list[dict[str, Any]],
    checks: list[dict[str, Any]],
) -> None:
    """确认定稿 CSV 的每个标识与 Sheet3 原始记录一一对应。"""
    source_by_sequence = {int(row["序号_整数"]): row for row in source_rows}
    matched_by_field = {column: 0 for column in ["序号", "关键词", "方案ID", "推广单元ID"]}
    complete_match_count = 0

    for _, score_row in scores.iterrows():
        sequence = int(score_row["序号_整数"])
        source_row = source_by_sequence.get(sequence)
        if source_row is None:
            continue
        row_matches = []
        for column in matched_by_field:
            matched = canonical_code(score_row[column]) == canonical_code(source_row[column])
            matched_by_field[column] += int(matched)
            row_matches.append(matched)
        complete_match_count += int(all(row_matches))

    for column, matched_count in matched_by_field.items():
        add_check(
            checks,
            f"score_source_{column}_matched_rows",
            EXPECTED_TOTAL,
            matched_count,
            matched_count == EXPECTED_TOTAL,
        )
    add_check(
        checks,
        "score_source_complete_matched_rows",
        EXPECTED_TOTAL,
        complete_match_count,
        complete_match_count == EXPECTED_TOTAL,
    )
    require_all_passed(checks, "评分表与 Sheet3 一一对应检查")


def template_snapshot(checks: list[dict[str, Any]]) -> dict[str, Any]:
    """读取真实模板结构和格式快照，严禁猜测模板。"""
    if not TEMPLATE_XLSX.exists():
        raise FileNotFoundError(f"未找到官方模板：{TEMPLATE_XLSX}")
    workbook = load_workbook(TEMPLATE_XLSX, data_only=False)
    sheet_names = workbook.sheetnames.copy()
    add_check(checks, "template_sheet_count", 1, len(sheet_names), len(sheet_names) == 1)
    add_check(
        checks,
        "template_sheet_name",
        "Sheet1",
        sheet_names,
        sheet_names == ["Sheet1"],
    )
    require_all_passed(checks, "模板工作表检查")

    sheet = workbook[sheet_names[0]]
    headers = [sheet.cell(1, column).value for column in range(1, 9)]
    add_check(
        checks,
        "template_header_order",
        EXPECTED_TEMPLATE_HEADERS,
        headers,
        headers == EXPECTED_TEMPLATE_HEADERS,
    )
    add_check(
        checks,
        "template_start_write_row",
        2,
        sheet.max_row + 1,
        sheet.max_row == 1,
    )
    require_all_passed(checks, "模板表头与写入位置检查")

    snapshot = {
        "sheet_names": sheet_names,
        "headers": headers,
        "header_style_ids": [sheet.cell(1, column).style_id for column in range(1, 9)],
        "column_widths": {
            key: dimension.width for key, dimension in sheet.column_dimensions.items()
        },
        "row_heights": {
            key: dimension.height for key, dimension in sheet.row_dimensions.items()
        },
        "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
        "freeze_panes": canonical_code(sheet.freeze_panes),
        "auto_filter": canonical_code(sheet.auto_filter.ref),
    }
    workbook.close()
    return snapshot


def write_result(
    scores: pd.DataFrame,
    source_rows: list[dict[str, Any]],
    snapshot: dict[str, Any],
) -> None:
    """基于原模板写出结果；每行只写一个类别关键词单元格。"""
    score_by_sequence = {
        int(row["序号_整数"]): row for _, row in scores.iterrows()
    }
    category_to_column = {
        category: EXPECTED_TEMPLATE_HEADERS.index(category) + 1
        for category in CATEGORY_ORDER
    }

    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = OUTPUT_XLSX.with_name(f".{OUTPUT_XLSX.stem}.tmp.xlsx")
    if temporary_path.exists():
        temporary_path.unlink()
    try:
        workbook = load_workbook(TEMPLATE_XLSX, data_only=False)
        sheet = workbook[snapshot["sheet_names"][0]]
        for output_row, source_row in enumerate(source_rows, start=2):
            sequence = int(source_row["序号_整数"])
            score_row = score_by_sequence[sequence]
            category = str(score_row["最终类别"])

            sheet.cell(output_row, 1).value = source_row["方案ID"]
            sheet.cell(output_row, 2).value = source_row["推广单元ID"]
            sheet.cell(output_row, 3).value = source_row["序号"]
            sheet.cell(output_row, category_to_column[category]).value = source_row["关键词"]

        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, OUTPUT_XLSX)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def validate_output(
    scores: pd.DataFrame,
    source_rows: list[dict[str, Any]],
    snapshot: dict[str, Any],
    checks: list[dict[str, Any]],
) -> dict[str, Any]:
    """重新读取结果文件，复核内容、唯一性、分类数量和模板一致性。"""
    workbook = load_workbook(OUTPUT_XLSX, read_only=False, data_only=False)
    output_sheet_names = workbook.sheetnames.copy()
    add_check(
        checks,
        "output_template_sheet_names",
        snapshot["sheet_names"],
        output_sheet_names,
        output_sheet_names == snapshot["sheet_names"],
    )
    sheet = workbook[output_sheet_names[0]]

    output_headers = [sheet.cell(1, column).value for column in range(1, 9)]
    header_styles = [sheet.cell(1, column).style_id for column in range(1, 9)]
    column_widths = {
        key: dimension.width for key, dimension in sheet.column_dimensions.items()
    }
    row_heights = {key: dimension.height for key, dimension in sheet.row_dimensions.items()}
    merged_ranges = [str(item) for item in sheet.merged_cells.ranges]
    add_check(
        checks,
        "output_template_headers",
        snapshot["headers"],
        output_headers,
        output_headers == snapshot["headers"],
    )
    add_check(
        checks,
        "output_template_header_styles",
        snapshot["header_style_ids"],
        header_styles,
        header_styles == snapshot["header_style_ids"],
    )
    add_check(
        checks,
        "output_template_column_widths",
        snapshot["column_widths"],
        column_widths,
        column_widths == snapshot["column_widths"],
    )
    add_check(
        checks,
        "output_template_row_heights",
        snapshot["row_heights"],
        row_heights,
        row_heights == snapshot["row_heights"],
    )
    add_check(
        checks,
        "output_template_merged_ranges",
        snapshot["merged_ranges"],
        merged_ranges,
        merged_ranges == snapshot["merged_ranges"],
    )
    output_freeze = canonical_code(sheet.freeze_panes)
    output_filter = canonical_code(sheet.auto_filter.ref)
    add_check(
        checks,
        "output_template_freeze_panes",
        snapshot["freeze_panes"],
        output_freeze,
        output_freeze == snapshot["freeze_panes"],
    )
    add_check(
        checks,
        "output_template_auto_filter",
        snapshot["auto_filter"],
        output_filter,
        output_filter == snapshot["auto_filter"],
    )

    output_rows: list[dict[str, Any]] = []
    for row_number in range(2, sheet.max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, 9)]
        if not any(is_nonempty(value) for value in values):
            continue
        category_values = values[3:8]
        populated_indices = [
            index for index, value in enumerate(category_values) if is_nonempty(value)
        ]
        output_rows.append(
            {
                "方案ID": values[0],
                "推广单元ID": values[1],
                "序号": values[2],
                "类别非空数": len(populated_indices),
                "输出类别": (
                    CATEGORY_ORDER[populated_indices[0]] if len(populated_indices) == 1 else None
                ),
                "关键词": (
                    category_values[populated_indices[0]]
                    if len(populated_indices) == 1
                    else None
                ),
            }
        )
    workbook.close()

    score_categories = scores["最终类别"].tolist()
    source_sequences = [row["序号"] for row in source_rows]
    source_keywords = [row["关键词"] for row in source_rows]
    source_plans = [row["方案ID"] for row in source_rows]
    source_units = [row["推广单元ID"] for row in source_rows]

    output_sequences = [row["序号"] for row in output_rows]
    output_keywords = [row["关键词"] for row in output_rows]
    output_plans = [row["方案ID"] for row in output_rows]
    output_units = [row["推广单元ID"] for row in output_rows]
    output_categories = [row["输出类别"] for row in output_rows]
    per_row_counts = [row["类别非空数"] for row in output_rows]

    duplicate_sequence_count = len(output_sequences) - len(
        {canonical_code(value) for value in output_sequences}
    )
    missing_classification_count = sum(count == 0 for count in per_row_counts)
    multiple_classification_count = sum(count > 1 for count in per_row_counts)
    exactly_one_count = sum(count == 1 for count in per_row_counts)

    add_check(
        checks,
        "output_record_count",
        EXPECTED_TOTAL,
        len(output_rows),
        len(output_rows) == EXPECTED_TOTAL,
    )
    add_check(
        checks,
        "output_duplicate_sequence_count",
        0,
        duplicate_sequence_count,
        duplicate_sequence_count == 0,
    )
    add_check(
        checks,
        "output_missing_classification_rows",
        0,
        missing_classification_count,
        missing_classification_count == 0,
    )
    add_check(
        checks,
        "output_multiple_classification_rows",
        0,
        multiple_classification_count,
        multiple_classification_count == 0,
    )
    add_check(
        checks,
        "output_exactly_one_category_rows",
        EXPECTED_TOTAL,
        exactly_one_count,
        exactly_one_count == EXPECTED_TOTAL,
    )

    actual_output_counts = {
        category: output_categories.count(category) for category in CATEGORY_ORDER
    }
    for category in CATEGORY_ORDER:
        expected = EXPECTED_COUNTS[category]
        actual = actual_output_counts[category]
        add_check(
            checks,
            f"output_category_count_{category}",
            expected,
            actual,
            actual == expected,
        )

    comparison_specs = [
        ("sequence", source_sequences, output_sequences),
        ("keyword", source_keywords, output_keywords),
        ("plan_id", source_plans, output_plans),
        ("unit_id", source_units, output_units),
        ("category", score_categories, output_categories),
    ]
    for name, expected_values, actual_values in comparison_specs:
        matched_count = sum(
            canonical_code(expected) == canonical_code(actual)
            for expected, actual in zip(expected_values, actual_values)
        )
        if len(expected_values) != len(actual_values):
            matched_count = min(matched_count, len(actual_values))
        add_check(
            checks,
            f"output_{name}_matched_rows",
            EXPECTED_TOTAL,
            matched_count,
            matched_count == EXPECTED_TOTAL
            and len(expected_values) == len(actual_values),
        )

    expected_records = {
        (
            canonical_code(row["方案ID"]),
            canonical_code(row["推广单元ID"]),
            canonical_code(row["序号"]),
            canonical_code(row["关键词"]),
        )
        for row in source_rows
    }
    output_records = {
        (
            canonical_code(row["方案ID"]),
            canonical_code(row["推广单元ID"]),
            canonical_code(row["序号"]),
            canonical_code(row["关键词"]),
        )
        for row in output_rows
    }
    add_check(
        checks,
        "output_unique_source_records",
        EXPECTED_TOTAL,
        len(output_records),
        len(output_records) == EXPECTED_TOTAL and output_records == expected_records,
    )

    return {
        "total": len(output_rows),
        "counts": actual_output_counts,
        "duplicate_sequences": duplicate_sequence_count,
        "missing_classifications": missing_classification_count,
        "multiple_classifications": multiple_classification_count,
    }


def print_final_summary(summary: dict[str, Any], all_passed: bool) -> None:
    """只打印用户指定的七项最终信息。"""
    print(f"1. result2.xlsx 输出位置：{OUTPUT_XLSX}")
    print(f"2. 总记录数：{summary['total']}")
    counts_text = "，".join(
        f"{category} {summary['counts'][category]}" for category in CATEGORY_ORDER
    )
    print(f"3. 五类数量：{counts_text}")
    print(f"4. 是否存在重复序号：{'是' if summary['duplicate_sequences'] else '否'}")
    print(f"5. 是否存在漏分类：{'是' if summary['missing_classifications'] else '否'}")
    print(f"6. 是否存在一行多类别：{'是' if summary['multiple_classifications'] else '否'}")
    print(f"7. 是否所有模板一致性检查通过：{'是' if all_passed else '否'}")


def main() -> None:
    checks: list[dict[str, Any]] = []
    scores = load_scores(checks)
    source_rows = load_source_rows(checks)
    validate_score_source_mapping(scores, source_rows, checks)
    snapshot = template_snapshot(checks)

    write_result(scores, source_rows, snapshot)
    summary = validate_output(scores, source_rows, snapshot, checks)
    validation = save_validation(checks)
    require_all_passed(checks, "输出文件回读校验")

    print_final_summary(summary, bool(validation["passed"].all()))


if __name__ == "__main__":
    main()
